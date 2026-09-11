import os
from collections import OrderedDict
from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.core.files import File
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction
from django.db.models import Q
from django.utils.text import slugify

from apps.products.models import Attribute, Brand, Category, Product, ProductImage, ProductVariant


PRODUCT_SPEC_MAPPINGS = OrderedDict(
    [
        (
            'Диагональ',
            {
                'slug': 'screen-size',
                'name': 'Диагональ экрана',
                'applies_to': 'product',
                'type': 'number',
                'unit': 'дюйм',
            },
        ),
        (
            'Связь',
            {
                'slug': 'connectivity',
                'name': 'Связь',
                'applies_to': 'product',
                'type': 'string',
            },
        ),
        (
            'Процессор',
            {
                'slug': 'processor',
                'name': 'Процессор',
                'applies_to': 'product',
                'type': 'string',
            },
        ),
        (
            'Год',
            {
                'slug': 'release-year',
                'name': 'Год выпуска',
                'applies_to': 'product',
                'type': 'number',
            },
        ),
        (
            'Разъём',
            {
                'slug': 'connector',
                'name': 'Разъём',
                'applies_to': 'product',
                'type': 'string',
            },
        ),
        (
            'Комплект',
            {
                'slug': 'bundle',
                'name': 'Комплект',
                'applies_to': 'product',
                'type': 'string',
            },
        ),
        (
            'MagSafe',
            {
                'slug': 'magsafe',
                'name': 'MagSafe',
                'applies_to': 'product',
                'type': 'enum',
            },
        ),
        (
            'Корпус',
            {
                'slug': 'body',
                'name': 'Корпус',
                'applies_to': 'product',
                'type': 'string',
            },
        ),
    ]
)

VARIANT_MAPPINGS = OrderedDict(
    [
        (
            'ОЗУ',
            {
                'slug': 'storage',
                'name': 'Память',
                'applies_to': 'variant',
                'type': 'enum',
            },
        ),
        (
            'Цвет_вариант',
            {
                'slug': 'color',
                'name': 'Цвет',
                'applies_to': 'variant',
                'type': 'enum',
            },
        ),
    ]
)

REQUIRED_COLUMNS = (
    'Серия',
    'Модель_из_файла',
    'ОЗУ',
    'Диагональ',
    'Цвет_вариант',
    'Цветы_всего',
    'Связь',
    'Процессор',
    'Год',
    'Разъём',
    'Комплект',
    'MagSafe',
    'Корпус',
    'gallery_id',
)

MEDIA_COLUMNS = ('Фото_1', 'Фото_2', 'Фото_3')


class DryRunRollback(Exception):
    pass


class Command(BaseCommand):
    help = 'Импортирует товары и варианты из XLSX файла.'

    def add_arguments(self, parser):
        parser.add_argument('xlsx_path', help='Путь к XLSX файлу')
        parser.add_argument('--category-slug', required=True, help='Slug категории товаров')
        parser.add_argument('--category-name', help='Название категории, если её нужно создать')
        parser.add_argument('--brand-slug', help='Slug бренда')
        parser.add_argument('--brand-name', help='Название бренда, если его нужно создать')
        parser.add_argument('--media-dir', help='Папка, в которой лежат файлы Фото_1..Фото_3')
        parser.add_argument(
            '--product-name-template',
            default='{brand} {model}',
            help='Шаблон имени товара. Доступны {brand}, {series}, {model}.',
        )
        parser.add_argument('--default-price', default='0', help='Цена по умолчанию для вариантов')
        parser.add_argument('--default-stock', type=int, default=0, help='Остаток по умолчанию для вариантов')
        parser.add_argument('--dry-run', action='store_true', help='Проверить импорт без сохранения')

    def handle(self, *args, **options):
        workbook = self._load_workbook(options['xlsx_path'])
        rows = self._extract_rows(workbook)
        if not rows:
            raise CommandError('В XLSX нет строк для импорта.')

        media_dir = Path(options['media_dir']).resolve() if options.get('media_dir') else None
        if media_dir and not media_dir.is_dir():
            raise CommandError(f'Папка с медиа не найдена: {media_dir}')

        category, _ = Category.objects.get_or_create(
            slug=options['category_slug'],
            defaults={'name': options.get('category_name') or options['category_slug']},
        )

        brand = None
        if options.get('brand_slug'):
            brand, _ = Brand.objects.get_or_create(
                slug=options['brand_slug'],
                defaults={'name': options.get('brand_name') or options['brand_slug']},
            )

        stats = {
            'products_created': 0,
            'products_updated': 0,
            'variants_created': 0,
            'variants_updated': 0,
            'media_created': 0,
            'media_updated': 0,
        }

        default_price = self._parse_decimal(options['default_price'], option_name='default-price')

        try:
            with transaction.atomic():
                self._ensure_attributes(category, self._collect_attribute_values(rows))

                for product_rows in self._group_rows_by_product(rows).values():
                    self._import_product_group(
                        category=category,
                        brand=brand,
                        product_rows=product_rows,
                        media_dir=media_dir,
                        product_name_template=options['product_name_template'],
                        default_price=default_price,
                        default_stock=options['default_stock'],
                        stats=stats,
                    )

                if options['dry_run']:
                    raise DryRunRollback
        except DryRunRollback:
            self.stdout.write(self.style.WARNING('Dry-run завершён, изменения не сохранены.'))
            self._print_stats(stats)
            return

        self.stdout.write(self.style.SUCCESS('Импорт завершён.'))
        self._print_stats(stats)

    def _load_workbook(self, xlsx_path):
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise CommandError(
                'Для импорта XLSX нужен пакет openpyxl. Установите зависимости проекта.'
            ) from exc

        xlsx_file = Path(xlsx_path).resolve()
        if not xlsx_file.is_file():
            raise CommandError(f'Файл не найден: {xlsx_file}')
        return load_workbook(filename=xlsx_file, read_only=True, data_only=True)

    def _extract_rows(self, workbook):
        worksheet = workbook.active
        iterator = worksheet.iter_rows(values_only=True)

        try:
            headers = [self._normalize_cell(value) for value in next(iterator)]
        except StopIteration as exc:
            raise CommandError('XLSX файл пуст.') from exc

        missing_columns = [column for column in REQUIRED_COLUMNS if column not in headers]
        if missing_columns:
            raise CommandError(f'В XLSX отсутствуют колонки: {", ".join(missing_columns)}')

        rows = []
        for row in iterator:
            row_data = {
                headers[index]: self._normalize_cell(value)
                for index, value in enumerate(row[: len(headers)])
            }
            if any(row_data.values()):
                rows.append(row_data)
        return rows

    def _collect_attribute_values(self, rows):
        values = {
            mapping['slug']: set()
            for mapping in [*PRODUCT_SPEC_MAPPINGS.values(), *VARIANT_MAPPINGS.values()]
        }

        for row in rows:
            for source_name, mapping in PRODUCT_SPEC_MAPPINGS.items():
                value = row.get(source_name)
                if value:
                    values[mapping['slug']].add(value)

            for source_name, mapping in VARIANT_MAPPINGS.items():
                value = row.get(source_name)
                if value:
                    values[mapping['slug']].add(value)

            for color in self._split_values(row.get('Цветы_всего', '')):
                values['color'].add(color)
        return values

    def _ensure_attributes(self, category, attribute_values):
        attributes = []
        for mapping in [*PRODUCT_SPEC_MAPPINGS.values(), *VARIANT_MAPPINGS.values()]:
            attribute = Attribute.objects.filter(
                Q(slug=mapping['slug']) | Q(name=mapping['name'])
            ).order_by('id').first()

            if attribute is None:
                attribute = Attribute.objects.create(
                    slug=mapping['slug'],
                    name=mapping['name'],
                    applies_to=mapping['applies_to'],
                    type=mapping['type'],
                    unit=mapping.get('unit', ''),
                    values=[],
                )

            update_fields = []
            if (
                attribute.slug != mapping['slug']
                and not Attribute.objects.exclude(pk=attribute.pk).filter(slug=mapping['slug']).exists()
            ):
                attribute.slug = mapping['slug']
                update_fields.append('slug')
            if attribute.name != mapping['name']:
                attribute.name = mapping['name']
                update_fields.append('name')
            if attribute.applies_to != mapping['applies_to']:
                attribute.applies_to = mapping['applies_to']
                update_fields.append('applies_to')
            if attribute.type != mapping['type']:
                attribute.type = mapping['type']
                update_fields.append('type')
            if attribute.unit != mapping.get('unit', ''):
                attribute.unit = mapping.get('unit', '')
                update_fields.append('unit')

            if mapping['type'] == 'enum':
                merged_values = list(
                    OrderedDict.fromkeys([*(attribute.values or []), *sorted(attribute_values[mapping['slug']])])
                )
                if list(attribute.values or []) != merged_values:
                    attribute.values = merged_values
                    update_fields.append('values')

            if update_fields:
                attribute.save(update_fields=update_fields)

            attributes.append(attribute)

        category.attributes.add(*attributes)

    def _group_rows_by_product(self, rows):
        grouped_rows = OrderedDict()
        for row in rows:
            product_key = (row.get('Серия', ''), row.get('Модель_из_файла', ''))
            grouped_rows.setdefault(product_key, []).append(row)
        return grouped_rows

    def _import_product_group(
        self,
        *,
        category,
        brand,
        product_rows,
        media_dir,
        product_name_template,
        default_price,
        default_stock,
        stats,
    ):
        first_row = product_rows[0]
        product_name = self._build_product_name(
            brand_name=brand.name if brand else '',
            series=first_row.get('Серия', ''),
            model=first_row.get('Модель_из_файла', ''),
            template=product_name_template,
        )
        product_slug = self._resolve_product_slug(product_name, category)
        specifications = self._build_product_specifications(first_row)

        product = Product.objects.filter(slug=product_slug).first()
        created = product is None
        if created:
            product = Product(slug=product_slug, category=category)

        product.name = product_name
        product.category = category
        product.brand = brand
        product.specifications = specifications
        product.short_description = product.short_description or product_name
        product.save()

        stats['products_created' if created else 'products_updated'] += 1

        for row in product_rows:
            self._upsert_variant(
                product=product,
                row=row,
                default_price=default_price,
                default_stock=default_stock,
                stats=stats,
            )

        if media_dir:
            self._sync_media(product=product, product_rows=product_rows, media_dir=media_dir, stats=stats)

    def _build_product_name(self, *, brand_name, series, model, template):
        name = template.format(
            brand=brand_name.strip(),
            series=series.strip(),
            model=(model or series).strip(),
        ).strip()
        return ' '.join(name.split())

    def _resolve_product_slug(self, product_name, category):
        existing = Product.objects.filter(category=category, name=product_name).first()
        if existing:
            return existing.slug

        base_slug = slugify(product_name) or slugify(category.slug) or 'product'
        if not Product.objects.filter(slug=base_slug).exists():
            return base_slug

        suffix = 2
        while Product.objects.filter(slug=f'{base_slug}-{suffix}').exists():
            suffix += 1
        return f'{base_slug}-{suffix}'

    def _build_product_specifications(self, row):
        specifications = {}
        for source_name, mapping in PRODUCT_SPEC_MAPPINGS.items():
            value = row.get(source_name)
            if value:
                specifications[mapping['slug']] = value
        return specifications

    def _upsert_variant(self, *, product, row, default_price, default_stock, stats):
        attributes = {}
        for source_name, mapping in VARIANT_MAPPINGS.items():
            value = row.get(source_name)
            if value:
                attributes[mapping['slug']] = value

        variant = None
        for existing_variant in product.variants.all():
            if existing_variant.attributes == attributes:
                variant = existing_variant
                break

        created = variant is None
        if created:
            variant = ProductVariant(product=product)

        variant.attributes = attributes
        if created or variant.price in (None, Decimal('0')):
            variant.price = default_price
        if created or variant.stock == 0:
            variant.stock = default_stock
        variant.is_active = True
        variant.save()

        stats['variants_created' if created else 'variants_updated'] += 1

    def _sync_media(self, *, product, product_rows, media_dir, stats):
        desired_media = OrderedDict()
        for row in product_rows:
            color_value = row.get('Цвет_вариант', '')
            for order, column_name in enumerate(MEDIA_COLUMNS):
                filename = row.get(column_name, '')
                if not filename:
                    continue
                media_meta = desired_media.setdefault(
                    filename,
                    {'colors': set(), 'order': order, 'color_hits': 0},
                )
                if color_value:
                    media_meta['colors'].add(color_value)
                    media_meta['color_hits'] += 1

        existing_media = {
            (os.path.basename(media.image.name), media.color_value or ''): media
            for media in product.images.all()
        }

        for filename, media_meta in desired_media.items():
            resolved_color = ''
            if media_meta['colors'] and media_meta['color_hits'] == 1 and len(media_meta['colors']) == 1:
                resolved_color = next(iter(media_meta['colors']))

            media_key = (filename, resolved_color)
            media = existing_media.get(media_key)
            created = media is None

            if created:
                file_path = media_dir / filename
                if not file_path.is_file():
                    self.stderr.write(self.style.WARNING(f'Файл медиа не найден: {file_path}'))
                    continue

                with file_path.open('rb') as file_handle:
                    media = ProductImage(
                        product=product,
                        color_value=resolved_color,
                        order=media_meta['order'],
                        alt_text=product.name,
                    )
                    media.image.save(filename, File(file_handle), save=False)
                    media.full_clean()
                    media.save()
            else:
                update_fields = []
                if media.order != media_meta['order']:
                    media.order = media_meta['order']
                    update_fields.append('order')
                if media.color_value != resolved_color:
                    media.color_value = resolved_color
                    update_fields.append('color_value')
                if not media.alt_text:
                    media.alt_text = product.name
                    update_fields.append('alt_text')
                if update_fields:
                    media.save(update_fields=update_fields)

            stats['media_created' if created else 'media_updated'] += 1

    def _parse_decimal(self, raw_value, *, option_name):
        try:
            return Decimal(str(raw_value))
        except (InvalidOperation, TypeError, ValueError) as exc:
            raise CommandError(f'Некорректное значение для --{option_name}: {raw_value}') from exc

    def _normalize_cell(self, value):
        if value is None:
            return ''
        if isinstance(value, str):
            return value.strip()
        return str(value).strip()

    def _split_values(self, raw_value):
        return [value.strip() for value in raw_value.split(';') if value and value.strip()]

    def _print_stats(self, stats):
        self.stdout.write(
            'Создано товаров: {products_created}, обновлено товаров: {products_updated}, '
            'создано вариантов: {variants_created}, обновлено вариантов: {variants_updated}, '
            'создано медиа: {media_created}, обновлено медиа: {media_updated}'.format(**stats)
        )
